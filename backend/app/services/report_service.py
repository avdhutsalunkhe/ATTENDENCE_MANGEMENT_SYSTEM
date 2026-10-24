from sqlalchemy.orm import Session
from datetime import datetime, date
import calendar
import io
from typing import Optional, Dict, Any, List

from app.models.employee import Employee
from app.models.attendance import DailyAttendance
from app.services.attendance_service import get_settings

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def generate_monthly_matrix(db: Session, year: int, month: int, department: Optional[str] = None) -> Dict[str, Any]:
    _, days_in_month = calendar.monthrange(year, month)
    start_dt = datetime(year, month, 1, 0, 0, 0)
    end_dt = datetime(year, month, days_in_month, 23, 59, 59)

    emp_query = db.query(Employee).filter(Employee.status == "Active")
    if department and department != "All":
        emp_query = emp_query.filter(Employee.department == department)
    employees = emp_query.order_by(Employee.name.asc()).all()

    # Query all daily attendances in this month
    emp_ids = [e.id for e in employees]
    attendances = db.query(DailyAttendance).filter(
        DailyAttendance.employee_id.in_(emp_ids),
        DailyAttendance.attendance_date >= start_dt,
        DailyAttendance.attendance_date <= end_dt
    ).all()

    att_map: Dict[str, Dict[int, DailyAttendance]] = {}
    for a in attendances:
        day_num = a.attendance_date.day
        att_map.setdefault(a.employee_id, {})[day_num] = a

    emp_list = []
    total_present_all = 0
    total_absent_all = 0
    total_halfday_all = 0
    total_hours_all = 0.0

    settings = get_settings(db)
    weekend_days = [w.strip().lower() for w in settings.get("weekend_days", "Saturday,Sunday").split(",") if w.strip()]

    for emp in employees:
        emp_records = att_map.get(emp.id, {})
        days_dict = {}
        pres_count = 0
        abs_count = 0
        hd_count = 0
        wo_count = 0
        leave_count = 0
        late_count = 0
        tot_mins = 0.0

        for day in range(1, days_in_month + 1):
            cur_date = date(year, month, day)
            day_name = cur_date.strftime("%A").lower()
            rec = emp_records.get(day)

            if rec:
                st = rec.attendance_status
                hrs = round(rec.credited_work_minutes / 60.0, 1)
                lt = rec.late_minutes
            else:
                if day_name in weekend_days:
                    st = "WO"
                else:
                    st = "—"
                hrs = 0.0
                lt = 0

            if st == "P":
                pres_count += 1
            elif st == "HD":
                hd_count += 1
            elif st == "A":
                abs_count += 1
            elif st == "WO":
                wo_count += 1
            elif st in ("LEAVE", "CL", "SL", "PL", "LWP"):
                leave_count += 1

            if lt > 0:
                late_count += 1

            tot_mins += (rec.credited_work_minutes if rec else 0.0)

            days_dict[str(day)] = {
                "status": st,
                "hours": hrs,
                "late": lt
            }

        workable_days = days_in_month - wo_count
        effective_present = pres_count + (0.5 * hd_count)
        rate = round((effective_present / workable_days * 100), 1) if workable_days > 0 else 0.0
        tot_hours = round(tot_mins / 60.0, 1)

        total_present_all += pres_count
        total_absent_all += abs_count
        total_halfday_all += hd_count
        total_hours_all += tot_hours

        emp_list.append({
            "id": emp.id,
            "name": emp.name,
            "department": emp.department or "Unassigned",
            "designation": emp.designation or "Staff",
            "days": days_dict,
            "totals": {
                "present": pres_count,
                "absent": abs_count,
                "half_day": hd_count,
                "weekly_off": wo_count,
                "leave": leave_count,
                "late_count": late_count,
                "total_hours": tot_hours,
                "attendance_rate": rate
            }
        })

    avg_rate = round(sum(e["totals"]["attendance_rate"] for e in emp_list) / len(emp_list), 1) if emp_list else 0.0

    return {
        "year": year,
        "month": month,
        "month_name": calendar.month_name[month],
        "days_in_month": days_in_month,
        "employees": emp_list,
        "summary": {
            "total_employees": len(emp_list),
            "average_attendance": avg_rate,
            "total_work_hours": round(total_hours_all, 1),
            "present_days": total_present_all,
            "absent_days": total_absent_all,
            "half_day_days": total_halfday_all
        }
    }


def generate_excel_report(db: Session, year: int, month: int, department: Optional[str] = None) -> io.BytesIO:
    matrix = generate_monthly_matrix(db, year, month, department)
    wb = openpyxl.Workbook()
    
    # Sheet 1: Monthly Summary
    ws_summary = wb.active
    ws_summary.title = "Monthly Summary"

    # Header styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    ws_summary.append([f"Attendance Monthly Summary - {matrix['month_name']} {year}"])
    ws_summary.cell(row=1, column=1).font = title_font
    ws_summary.append([f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Department: {department or 'All'}"])
    ws_summary.cell(row=2, column=1).font = subtitle_font
    ws_summary.append([]) # empty row

    headers = [
        "Employee ID", "Full Name", "Department", "Designation",
        "Present Days", "Half Days", "Absent Days", "Weekly Offs",
        "Late Incidents", "Total Work Hours", "Attendance %"
    ]
    ws_summary.append(headers)

    h_row = 4
    for col_idx in range(1, len(headers) + 1):
        cell = ws_summary.cell(row=h_row, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx > 4 else "left", vertical="center")

    row_num = 5
    for emp in matrix["employees"]:
        t = emp["totals"]
        row_data = [
            emp["id"], emp["name"], emp["department"], emp["designation"],
            t["present"], t["half_day"], t["absent"], t["weekly_off"],
            t["late_count"], t["total_hours"], f"{t['attendance_rate']}%"
        ]
        ws_summary.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws_summary.cell(row=row_num, column=col_idx)
            cell.border = border_thin
            if col_idx > 4:
                cell.alignment = Alignment(horizontal="center")
        row_num += 1

    # Adjust widths
    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Sheet 2: Daily Attendance Matrix
    ws_matrix = wb.create_sheet(title="Daily Matrix")
    ws_matrix.append([f"Daily Attendance Grid - {matrix['month_name']} {year}"])
    ws_matrix.cell(row=1, column=1).font = title_font
    ws_matrix.append([])

    matrix_headers = ["ID", "Name", "Department"] + [str(d) for d in range(1, matrix["days_in_month"] + 1)] + ["Total Hrs", "Rate"]
    ws_matrix.append(matrix_headers)

    for col_idx in range(1, len(matrix_headers) + 1):
        cell = ws_matrix.cell(row=3, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 4
    fill_present = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # green
    fill_absent = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # red
    fill_half = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")    # yellow
    fill_wo = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")      # gray

    for emp in matrix["employees"]:
        row_data = [emp["id"], emp["name"], emp["department"]]
        day_statuses = [emp["days"].get(str(d), {}).get("status", "—") for d in range(1, matrix["days_in_month"] + 1)]
        row_data.extend(day_statuses)
        row_data.append(emp["totals"]["total_hours"])
        row_data.append(f"{emp['totals']['attendance_rate']}%")
        ws_matrix.append(row_data)

        # Style status cells
        for idx, st in enumerate(day_statuses, start=4):
            cell = ws_matrix.cell(row=row_num, column=idx)
            cell.alignment = Alignment(horizontal="center")
            cell.border = border_thin
            if st == "P":
                cell.fill = fill_present
            elif st == "A":
                cell.fill = fill_absent
            elif st == "HD":
                cell.fill = fill_half
            elif st == "WO":
                cell.fill = fill_wo

        row_num += 1

    # Format matrix widths
    ws_matrix.column_dimensions["A"].width = 12
    ws_matrix.column_dimensions["B"].width = 22
    ws_matrix.column_dimensions["C"].width = 16
    for d in range(1, matrix["days_in_month"] + 1):
        ws_matrix.column_dimensions[get_column_letter(3 + d)].width = 5

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def generate_pdf_report(db: Session, year: int, month: int, department: Optional[str] = None) -> io.BytesIO:
    matrix = generate_monthly_matrix(db, year, month, department)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'RepTitle', parent=styles['Heading1'],
        fontSize=18, textColor=colors.HexColor('#0F172A'), spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'RepSub', parent=styles['Normal'],
        fontSize=9, textColor=colors.HexColor('#64748B'), spaceAfter=14
    )

    elements = []
    elements.append(Paragraph(f"Monthly Attendance Report: {matrix['month_name']} {year}", title_style))
    elements.append(Paragraph(
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Department: {department or 'All'} | Staff Count: {matrix['summary']['total_employees']} | Overall Attendance: {matrix['summary']['average_attendance']}%",
        subtitle_style
    ))

    # Summary table
    table_data = [
        ["Emp ID", "Employee Name", "Department", "Designation", "Present", "Half-Day", "Absent", "Late", "Total Hours", "Att. Rate"]
    ]

    for emp in matrix["employees"]:
        t = emp["totals"]
        table_data.append([
            emp["id"],
            emp["name"],
            emp["department"],
            emp["designation"],
            str(t["present"]),
            str(t["half_day"]),
            str(t["absent"]),
            str(t["late_count"]),
            f"{t['total_hours']}h",
            f"{t['attendance_rate']}%"
        ])

    pdf_table = Table(table_data, colWidths=[65, 130, 95, 95, 55, 55, 55, 45, 75, 65])
    pdf_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (1, 1), (3, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 1), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 4),
    ]))

    elements.append(pdf_table)
    doc.build(elements)
    buffer.seek(0)
    return buffer
